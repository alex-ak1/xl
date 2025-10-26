import openpyxl
from dataclasses import dataclass
from jsonpointer import resolve_pointer, JsonPointerException
from io import BytesIO

@dataclass
class ColumnDef:
    id: int
    caption: str
    path: str # for fetching value from each data row

def makeExcel(data:list, columns:list[ColumnDef]) -> BytesIO:
    wb = openpyxl.Workbook()
    sheet = wb.active

    for i in range(len(columns)):
        sheet.cell(column=i+1, row=1, value=columns[i].caption)

    crow = 2
    for row in data:
        for i in range(len(columns)):
            try:
                v = resolve_pointer(row, columns[i].path)
            except JsonPointerException:
                continue

            sheet.cell(column=i+1, row=crow, value=v)

        crow += 1

    r = BytesIO()
    wb.save(r)
    r.seek(0)

    return r





