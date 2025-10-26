import pytest
from makeXl import makeExcel, ColumnDef
import openpyxl
import os

fname = "test.xlsx"
def make_test_sheet():
    cols = [
        ColumnDef(1, "Название", "/addr"),
        ColumnDef(2, "поле 1", "/sites/1/value"),
        ColumnDef(3, "поле 2", "/sites/2/value"),
        ColumnDef(4, "поле 3", "/sites/3/value"),
        ColumnDef(5, "поле 4", "/sites/4/value"),
        ColumnDef(6, "поле 5", "/sites/5/value"),
    ]

    data = [
        { "addr": "ya.ru",
          "sites": {
              "1": {"value": 121},
              "2": {"value": 122},
              "3": {"value": 123}
          }},
        {"addr": "не ya.ru",
         "sites": {
             "1": {"value": 221},
             "2": {"value": 222},
             "4": {"value": 224}
         }},
        {"addr": "rambler.ru",
         "sites": {
             "1": {"value": 321},
             "4": {"value": 324},
             "5": {"value": 325}
         }}
    ]

    r = makeExcel(data, cols)
    with open(fname, "wb") as f:
        f.write(r.getbuffer())

def test_xl():
    if os.path.exists(fname):
        os.remove(fname)

    assert not os.path.exists(fname), "Cant delete file"

    make_test_sheet()

    wb = openpyxl.load_workbook(fname)
    sh = wb.active

    assert sh.cell(row=2, column=2).value == 121
    assert sh.cell(row=4, column=4).value is None